"""Compact identity metadata and bounded, lazy pixel access.

No optional training dependency is imported while opening ordinary image sources.
The index stores columns, not one Python Sample/ImageRef object per observation.
"""
from __future__ import annotations

from collections import OrderedDict
from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import Path
import hashlib
import json
import math
import re
import string

import numpy as np
import pandas as pd
from PIL import Image

from ..models import Annotation, DatasetMetadata, Sample, Task
from ..utils import normalize_split


@dataclass(frozen=True)
class ImageRef:
    path: str
    array: str | None = None
    row: int | None = None
    crop: tuple[int, int, int, int] | None = None
    trim_black_borders: bool = False

    @property
    def key(self):
        values = (self.path, self.array, self.row, self.crop)
        return json.dumps((*values, True) if self.trim_black_borders else values, separators=(",", ":"))

    def read(self):
        if self.array is None:
            with Image.open(self.path) as image:
                pixels = np.array(image.convert("RGB"))
        else:
            pixels = np.array(_reader(self.path)[self.array][self.row], copy=True)
            if pixels.ndim == 2:
                pixels = np.repeat(pixels[..., None], 3, axis=2)
            if pixels.shape[-1] == 1:
                pixels = np.repeat(pixels, 3, axis=2)
            if pixels.ndim != 3 or pixels.shape[-1] != 3 or pixels.dtype != np.uint8:
                raise ValueError("Identity images must be H×W×3 uint8 (or grayscale)")
        if self.crop is not None:
            x0, y0, x1, y1 = self.crop
            pixels = pixels[y0:y1, x0:x1]
        if self.trim_black_borders:
            foreground = np.any(pixels != 0, axis=2)
            rows = np.flatnonzero(foreground.any(axis=1))
            columns = np.flatnonzero(foreground.any(axis=0))
            if len(rows) and len(columns):
                pixels = pixels[rows[0]:rows[-1]+1, columns[0]:columns[-1]+1]
        return np.ascontiguousarray(pixels)


# Process-local LRU: bounded archive mappings, never a decoded image cache.
# Diverse batches revisit many clips; two entries repeatedly reopen the archives.
_READER_LIMIT = 32
_READERS = OrderedDict()
_READER_PID = None


def _reader(path):
    import os
    global _READER_PID
    if _READER_PID != os.getpid():
        _READERS.clear()
        _READER_PID = os.getpid()
    if path not in _READERS:
        try:
            from mmnpz import NpzReader
            # mmnpz 1.0 uses two NumPy header readers moved in NumPy 2.3.
            # Adapt only mmnpz's module-local NumPy reference; never patch NumPy.
            if not hasattr(np.lib.format, "_read_array_header"):
                from types import SimpleNamespace
                import mmnpz.reading
                from numpy.lib import _format_impl
                local_numpy = SimpleNamespace(**vars(np))
                local_numpy.lib = SimpleNamespace(**vars(np.lib))
                local_numpy.lib.format = _format_impl
                mmnpz.reading.np = local_numpy
        except ImportError as error:
            raise ImportError("NPZ identity sources require pip install 'dataset-fixer[id]'") from error
        _READERS[path] = NpzReader(path)
        while len(_READERS) > _READER_LIMIT:
            _, reader = _READERS.popitem(last=False)
            reader.close()
    _READERS.move_to_end(path)
    return _READERS[path]


_DEFAULTS = dict(array=None, row=None, crop=None, identity_id=None, clip_id=None,
                 frame_index=None, local_id=None, tracklet_id=None, split="train", notes=None,
                 original_id=None, trim_black_borders=False)


class SampleIndex(Sequence):
    """Immutable-view column index; Samples and ImageRefs are created on access."""
    def __init__(self, records, *, names=None, sampling=None):
        frame = records.copy() if isinstance(records, pd.DataFrame) else pd.DataFrame(records)
        for key, value in _DEFAULTS.items():
            if key not in frame:
                frame[key] = value
        if "path" not in frame or not {"width", "height"} <= set(frame):
            raise ValueError("Identity records need path, width and height")
        frame = frame.reset_index(drop=True)
        for key in ("path", "array", "identity_id", "clip_id", "tracklet_id", "split", "notes"):
            frame[key] = frame[key].astype("category")
        self.frame = frame
        self.names = names or {i: label for i, label in enumerate(sorted(frame.identity_id.dropna().unique()))}
        self.labels = {v: k for k, v in self.names.items()}
        self.sampling = dict(sampling or {})

    def __len__(self):
        return len(self.frame)

    def __getitem__(self, index):
        if isinstance(index, slice):
            return self.take(np.arange(len(self))[index])
        row = self.frame.iloc[int(index)]
        def clean(value):
            if value is None or (not isinstance(value, (tuple, list, dict, np.ndarray)) and pd.isna(value)):
                return None
            return value.item() if isinstance(value, np.generic) else value
        ref = ImageRef(str(row.path), clean(row["array"]), None if pd.isna(row.row) else int(row.row), clean(row.crop),
                       bool(clean(row.trim_black_borders)))
        identity = clean(row.identity_id)
        provenance = {"original_id": clean(row.original_id) or ref.key, "identity_notes": clean(row.notes), "tracklet_id": clean(row.tracklet_id), "clip_id": clean(row.clip_id), "identity_id": identity, "image_ref": ref.key}
        return Sample(Path(ref.path), Path(f"{hashlib.sha256(ref.key.encode()).hexdigest()[:20]}-{Path(ref.path).name}"), str(row.split),
                      int(row.width), int(row.height),
                      [Annotation(self.labels[identity], identity_id=identity)] if identity is not None else [],
                      provenance=provenance, identity_id=identity, clip_id=clean(row.clip_id),
                      frame_index=clean(row.frame_index), local_id=clean(row.local_id), image_ref=ref)

    def take(self, rows):
        return SampleIndex(self.frame.iloc[rows], names=self.names, sampling=self.sampling)

    def where(self, **conditions):
        mask = np.ones(len(self), dtype=bool)
        for key, values in conditions.items():
            values = [values] if isinstance(values, (str, int)) else values
            mask &= self.frame[key].isin(values).to_numpy()
        return self.take(np.flatnonzero(mask))

    @property
    def splits(self):
        return tuple(s for s in ("train", "val", "test") if s in self.frame.split.values)

    @property
    def fingerprint(self):
        # Metadata changes, source replacements, derived crops and role changes invalidate cohorts.
        frame = self.frame.astype(object).copy()
        frame["crop"] = frame.crop.map(str)
        stamps = [(p, Path(p).stat().st_size, Path(p).stat().st_mtime_ns) for p in sorted(frame.path.unique())]
        digest = hashlib.sha256(json.dumps(stamps).encode())
        digest.update(pd.util.hash_pandas_object(frame, index=False).values.tobytes())
        return digest.hexdigest()

    def weights(self, balance_by=None):
        fields = tuple(balance_by or self.sampling.get("balance_by", ("identity_id",)))
        if not fields or any(field not in self.frame for field in fields):
            raise ValueError("balance_by must contain metadata column names")
        if self.frame[list(fields)].isna().any().any():
            raise ValueError(f"Sampling requires non-missing {fields}; choose only available metadata")
        weights = np.ones(len(self), dtype=np.float64)
        # Hierarchical probability: equal identities, then equal children at every level.
        for depth in range(1, len(fields)):
            prefix = list(fields[:depth])
            counts = self.frame.groupby(prefix, observed=True)[fields[depth]].transform("nunique")
            weights /= counts.to_numpy()
        counts = self.frame.groupby(list(fields), observed=True)[fields[0]].transform("size")
        weights /= counts.to_numpy()
        return weights


def index_of(dataset):
    if isinstance(dataset._samples, SampleIndex):
        return dataset._samples
    if dataset.task != Task.ID:
        raise ValueError("Use id.from_dataset(..., identity=...) to explicitly select individual labels")
    return SampleIndex([record_from_sample(s) for s in dataset._samples], names=dataset.classes)


def record_from_sample(sample):
    ref = sample.image_ref or ImageRef(str(sample.image_path))
    return dict(path=ref.path, array=ref.array, row=ref.row, crop=ref.crop, trim_black_borders=ref.trim_black_borders,
                width=sample.width, height=sample.height, identity_id=sample.identity_id,
                clip_id=sample.clip_id, frame_index=sample.frame_index, local_id=sample.local_id,
                tracklet_id=sample.provenance.get("tracklet_id"), split=sample.split,
                original_id=sample.provenance.get("original_id", ref.key))


def as_dataset(index, *, parent=None, location=None, name=None, source_format="identity_manifest", warnings=()):
    from ..dataset import Dataset
    metadata = parent._metadata.copy() if parent else DatasetMetadata(names=index.names)
    return Dataset(location=parent.location if parent else Path(location), name=name or (parent.name if parent else Path(location).stem),
                   task=Task.ID, metadata=metadata, samples=index,
                   manifest=dict(parent._manifest) if parent else {}, data_yaml=None,
                   source_format=parent.format if parent else source_format,
                   warnings=[*(parent.warnings if parent else ()), *warnings], provenance={},
                   planned_splits=index.splits)


def open_dataset(location, *, format, layout=None, identities=None, name=None, trim_black_borders=None):
    if trim_black_borders is not None and not isinstance(trim_black_borders, bool):
        raise TypeError("trim_black_borders must be a boolean or None")
    path = Path(location).expanduser().resolve()
    if format == "identity_folders":
        if not layout:
            raise ValueError("identity_folders requires an explicit layout, e.g. '{split}/{identity_id}/{clip_id}/{image}'")
        parts = []
        fields = set()
        for literal, field, spec, conversion in string.Formatter().parse(layout):
            parts.append(re.escape(literal))
            if field:
                if field not in {"split", "identity_id", "clip_id", "image", "tracklet_id"} or field in fields:
                    raise ValueError(f"Unsupported or repeated layout field: {field}")
                fields.add(field)
                parts.append(f"(?P<{field}>[^/]+)")
        if not {"image", "identity_id"} <= fields:
            raise ValueError("layout requires image and identity_id fields")
        pattern = re.compile("".join(parts))
        records = []
        for file in sorted(path.rglob("*")):
            if file.is_file() and file.suffix.lower() in {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp"}:
                match = pattern.fullmatch(file.relative_to(path).as_posix())
                if match is None:
                    raise ValueError(f"Image does not match layout: {file}")
                records.append(_file_record({**match.groupdict(), "path": str(file)}, path))
        frame = pd.DataFrame(records)
    elif format == "identity_manifest":
        if path.suffix.lower() == ".csv":
            records = pd.read_csv(path, dtype={"identity_id": "string", "clip_id": "string"}).to_dict("records")
        elif path.suffix.lower() == ".jsonl":
            records = [json.loads(line) for line in path.read_text().splitlines() if line.strip()]
        else:
            raise ValueError("Identity manifests must be CSV or JSONL")
        frame = pd.DataFrame([_file_record(record, path.parent) for record in records])
    elif format == "trex_npz":
        frame = _trex_records(path, identities)
    else:
        raise ValueError(f"Unknown identity format: {format}")
    if frame.empty:
        raise ValueError("No identity images found")
    if trim_black_borders is not None or "trim_black_borders" not in frame:
        frame["trim_black_borders"] = True if trim_black_borders is None else trim_black_borders
    index = SampleIndex(frame)
    unlabeled = int(index.frame.identity_id.isna().sum())
    warnings = [f"{unlabeled} unlabeled crops retained; excluded from supervised training, not treated as unknown"] if unlabeled else []
    warnings.extend(frame.attrs.get("warnings", ()))
    result = as_dataset(index, location=path if path.is_dir() else path.parent, name=name or path.stem, source_format=format, warnings=warnings)
    result._metadata.extra["tracklets"] = frame.attrs.get("tracklets", {})
    return result


def _file_record(record, root):
    record = {k: (None if np.isscalar(v) and pd.isna(v) else v) for k, v in record.items()}
    archive = record.pop("archive", None)
    record.setdefault("trim_black_borders", True)
    if not isinstance(record["trim_black_borders"], (bool, np.bool_)):
        raise TypeError("Manifest trim_black_borders must be a boolean")
    source = archive or record.pop("image_path", None) or record.get("path")
    if not source:
        raise ValueError("Manifest records require image_path/path or archive, array, row")
    path = Path(source).expanduser()
    path = path.resolve() if path.is_absolute() else (root / path).resolve()
    record["path"] = str(path)
    record["split"] = normalize_split(record.get("split") or "train")
    if record.get("identity_id") is not None:
        record["identity_id"] = str(record["identity_id"])
    if archive or record.get("array"):
        if record.get("row") is None or not record.get("array"):
            raise ValueError("An archive reference requires array and row")
        array = _reader(str(path))[record["array"]]
        row = int(record["row"])
        if len(array.shape) not in (3, 4) or not 0 <= row < len(array):
            raise ValueError(f"Invalid image array/row: {path}, {row}")
        record.update(row=row, width=array.shape[2], height=array.shape[1])
    else:
        with Image.open(path) as image:
            record.update(width=image.width, height=image.height)
    if record.get("crop") is not None:
        crop = json.loads(record["crop"]) if isinstance(record["crop"], str) else record["crop"]
        record["crop"] = tuple(map(int, crop))
        x0, y0, x1, y1 = record["crop"]
        if not (0 <= x0 < x1 <= record["width"] and 0 <= y0 < y1 <= record["height"]):
            raise ValueError("Crop must lie within the source image")
        record.update(width=x1-x0, height=y1-y0)
    return record


def _identity_mapping(path):
    if path is None:
        return {}
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Global ID Summary"]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows)
        result = {}
        for row in rows:
            identity = row[0]
            if identity is None:
                continue
            for clip, value in zip(headers[1:], row[1:]):
                if value is None or not clip or not re.match(r"\d+_\d+_", str(clip)):
                    continue
                for observation, local in re.findall(r"(\d+)_ID(\d+)", str(value)):
                    if int(observation) != int(str(clip).split("_")[1]):
                        raise ValueError(f"Workbook observation mismatch: {clip}: {value}")
                    key = (str(clip), int(local))
                    if key in result and result[key][0] != str(identity):
                        raise ValueError(f"Conflicting identity assignments for {key}")
                    result[key] = (str(identity), str(value))
        return result
    finally:
        workbook.close()


def _trex_records(root, identities):
    mapping = _identity_mapping(identities)
    frames, tables, warnings = [], {}, []
    paths = sorted(root.rglob("*.npz")) if root.is_dir() else [root]
    for path in paths:
        reader = _reader(str(path))
        images = reader["images"]
        if images.ndim != 4 or images.shape[-1] not in (1, 3):
            raise ValueError("TRex identity import requires fixed-size NHWC crops")
        n = len(images)
        ids, numbers = np.asarray(reader["ids"]), np.asarray(reader["frames"])
        if ids.shape != (n,) or numbers.shape != (n,):
            raise ValueError(f"TRex row metadata does not match images: {path}")
        clip = path.parent.name
        labels = {int(local): mapping.get((clip, int(local)), (None, None)) for local in np.unique(ids)}
        frame = pd.DataFrame(dict(path=str(path), array="images", row=np.arange(n),
            width=images.shape[2], height=images.shape[1], clip_id=clip, local_id=ids,
            frame_index=numbers, identity_id=pd.Series(ids).map({k: v[0] for k, v in labels.items()}),
            notes=pd.Series(ids).map({k: v[1] for k, v in labels.items()})))
        if "frame_segment_indexes" in reader and "tracklets" in reader:
            segments = np.asarray(reader["frame_segment_indexes"])
            bounds = np.asarray(reader["tracklets"])
            if segments.shape != (n,) or bounds.ndim != 2 or bounds.shape[1] != 2:
                raise ValueError(f"Invalid TRex tracklet metadata: {path}")
            # Export parts retain GLOBAL segment indexes but reset the bounds
            # table, sometimes mid-tracklet. Never index that local table with a
            # global ID or infer boundaries from gaps between sampled frames.
            frame["frame_segment_index"] = segments
            frame["tracklet_id"] = pd.Series(ids).astype(str) + ":" + pd.Series(segments).astype(str)
            tables[str(path)] = bounds.tolist()
            if segments.max(initial=0) >= len(bounds):
                warnings.append(f"{path.name}: global tracklet indexes exceed the part-local bounds table; native segment groups retained, bounds preserved separately")
        for key in ("path", "array", "clip_id", "identity_id", "notes", "tracklet_id"):
            if key in frame:
                frame[key] = frame[key].astype("category")
        frames.append(frame)
    combined = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()
    combined.attrs.update(tracklets=tables, warnings=warnings)
    return combined


def from_dataset(dataset, *, identity, context_fraction=0.0):
    """Create lazy crops before custom augmentations; never resize/rotate here.

    identity='class' explicitly treats annotation classes as individuals.
    identity='identity_id' uses independent labels, or supply (sample, annotation)->label.

    Parameters:
        dataset: Existing annotated Dataset to crop lazily.
        identity: Explicit class/identity_id selector or identity-label callable.
        context_fraction: Extra box-width/height fraction on each side, clipped to source bounds.
    """
    if not math.isfinite(context_fraction) or context_fraction < 0:
        raise ValueError("context_fraction must be finite and non-negative")
    records = []
    for sample in dataset._samples:
        for annotation in sample.annotations:
            label = dataset.classes[annotation.class_id] if identity == "class" else (
                annotation.identity_id if identity == "identity_id" else identity(sample, annotation) if callable(identity) else None)
            if identity not in ("class", "identity_id") and not callable(identity):
                raise ValueError("identity must be 'class', 'identity_id', or a callable")
            box = annotation.bbox or (0, 0, sample.width, sample.height)
            x0, y0, x1, y1 = box
            dx, dy = (x1-x0)*context_fraction, (y1-y0)*context_fraction
            crop = (max(0, math.floor(x0-dx)), max(0, math.floor(y0-dy)),
                    min(sample.width, math.ceil(x1+dx)), min(sample.height, math.ceil(y1+dy)))
            if crop[2] <= crop[0] or crop[3] <= crop[1]:
                raise ValueError("Cannot crop an empty bounding box")
            record = record_from_sample(sample)
            # Compose with an existing derived crop in original pixel coordinates.
            if record["crop"]:
                left, top, _, _ = record["crop"]
                crop = (crop[0]+left, crop[1]+top, crop[2]+left, crop[3]+top)
            record.update(crop=crop, identity_id=None if label is None else str(label),
                          width=crop[2]-crop[0], height=crop[3]-crop[1])
            records.append(record)
    return as_dataset(SampleIndex(records), location=dataset.location, name=dataset.name)


def split_dataset(dataset, ratios, *, group_by=None, assign=None, source_splits=None, seed=42, name=None,
                  classes_must_appear_in=()):
    from ..planning import allocate_groups, normalize_split_ratios
    index = index_of(dataset)
    if source_splits:
        index = index.where(split=[normalize_split(s) for s in source_splits])
    frame = index.frame.copy()
    ratios = normalize_split_ratios(ratios)
    if callable(group_by):
        groups = frame.path.map(lambda p: group_by(Path(p)))
    elif group_by:
        fields = [group_by] if isinstance(group_by, str) else list(group_by)
        if frame[fields].isna().any().any():
            raise ValueError(f"Grouping requires non-missing {fields}")
        groups = pd.Series(list(frame[fields].itertuples(index=False, name=None)))
    else:
        groups = pd.Series(np.arange(len(frame)))
    sizes = groups.value_counts(sort=False).to_dict()
    locked = {}
    if assign:
        for group, path in zip(groups, frame.path):
            target = assign(Path(path))
            if target is not None:
                target = normalize_split(target)
                if target not in ratios or (group in locked and locked[group] != target):
                    raise ValueError("Conflicting/invalid explicit group assignment")
                locked[group] = target
    group_classes = None
    if classes_must_appear_in:
        pairs = pd.DataFrame({"group": groups, "identity_id": frame.identity_id}).dropna(subset=["identity_id"])
        group_classes = {}
        for (group, label), count in pairs.groupby(["group", "identity_id"], sort=False, observed=True).size().items():
            group_classes.setdefault(group, {})[label] = int(count)
    allocated = allocate_groups(sizes, ratios, locked, seed, group_classes=group_classes,
                                classes_must_appear_in=classes_must_appear_in)
    frame["split"] = groups.map(allocated)
    warnings = []
    if classes_must_appear_in:
        observed = set(frame.identity_id.dropna())
        for split, ratio in ratios.items():
            absent = sorted(observed - set(frame.loc[frame.split == split, "identity_id"].dropna()))
            if ratio > 0 and absent:
                warnings.append(f"Split {split!r} lacks coverage for {len(absent)} identities: {', '.join(absent[:20])}" +
                                (" …" if len(absent) > 20 else "") + "; group boundaries preserved")
    result = as_dataset(SampleIndex(frame, names=index.names, sampling=index.sampling), parent=dataset, name=name, warnings=warnings)
    result._manifest["split_policy"] = dict(group_by=str(group_by), seed=seed, ratios=ratios)
    if classes_must_appear_in:
        result._manifest["split_policy"]["classes_must_appear_in"] = list(classes_must_appear_in)
    return result


def holdout(dataset, *, clip_id, identity_ids=None, to_split="test"):
    index = index_of(dataset)
    frame = index.frame.copy()
    mask = frame.clip_id == clip_id
    if identity_ids is not None:
        mask &= frame.identity_id.isin(identity_ids)
    if not mask.any():
        raise ValueError("No observations match this holdout")
    frame["split"] = frame.split.astype(object)
    frame.loc[mask, "split"] = normalize_split(to_split)
    shared = frame.loc[(frame.clip_id == clip_id) & (frame.split == "train")]
    warnings = [f"Clip {clip_id} remains in training for other individuals: held-out crops share frames/backgrounds"] if len(shared) and to_split != "train" else []
    return as_dataset(SampleIndex(frame, names=index.names, sampling=index.sampling), parent=dataset, warnings=warnings)


def with_sampling(dataset, *, balance_by=("identity_id", "clip_id"), samples_per_epoch=None, max_per_batch=None):
    """Balance weighted draws, then spread clips/tracklets within full batches.

    Batch ordering is best effort with eight batches of metadata lookahead.
    No draws are dropped or replaced. Optional max_per_batch caps remain strict
    and can produce smaller batches; no caps are applied by default.
    """
    index = index_of(dataset)
    if samples_per_epoch is not None and (isinstance(samples_per_epoch, bool) or not isinstance(samples_per_epoch, int) or samples_per_epoch < 1):
        raise ValueError("samples_per_epoch must be a positive integer")
    caps = dict(max_per_batch or {})
    for key, value in caps.items():
        if key not in {"identity_id", "clip_id", "tracklet_id"} or not isinstance(value, int) or value < 1:
            raise ValueError("Batch caps must map identity_id, clip_id or tracklet_id to positive integers")
    sampling = dict(balance_by=tuple(balance_by), samples_per_epoch=samples_per_epoch, max_per_batch=caps)
    selected = index.where(split="train")
    selected = selected.take(np.flatnonzero(selected.frame.identity_id.notna()))
    selected.weights(balance_by)  # Validate before training starts.
    if caps and selected.frame[list(caps)].isna().any().any():
        raise ValueError("Batch diversity limits require complete grouping metadata")
    return as_dataset(SampleIndex(index.frame, names=index.names, sampling=sampling), parent=dataset)


def balanced_rows(index, n, seed=7):
    """Bounded deterministic preview/gallery selection, round-robin across groups."""
    rng = np.random.default_rng(seed)
    frame = index.frame
    fields = [k for k in ("identity_id", "clip_id", "tracklet_id") if frame[k].notna().all()]
    def interleave(rows, depth):
        if depth == len(fields):
            yield from rng.choice(rows, size=min(n, len(rows)), replace=False)
            return
        groups = list(frame.iloc[rows].groupby(fields[depth], observed=True, sort=True).indices.values())
        rng.shuffle(groups)
        streams = [iter(interleave(rows[group], depth+1)) for group in groups]
        while streams:
            remaining = []
            for stream in streams:
                try:
                    yield next(stream)
                    remaining.append(stream)
                except StopIteration:
                    pass
            streams = remaining
    import itertools
    return list(itertools.islice(interleave(np.arange(len(index)), 0), n))



def export_manifest(dataset, destination):
    """Persist an immutable view as references; image/NPZ pixels are never copied."""
    destination = Path(destination)
    if destination.suffix != ".jsonl":
        destination = destination / "samples.jsonl"
    destination.parent.mkdir(parents=True, exist_ok=True)
    index = index_of(dataset)
    if destination.exists():
        raise FileExistsError(f"Refusing to overwrite {destination}")
    with destination.open("x") as handle:
        for values in index.frame.itertuples(index=False, name=None):
            record = dict(zip(index.frame.columns, values))
            record = {k: (None if v is None or (np.isscalar(v) and pd.isna(v)) else v) for k, v in record.items()}
            handle.write(json.dumps(record) + "\n")
    return open_dataset(destination, format="identity_manifest", name=dataset.name)


def summary_counts(dataset):
    """Metadata-only adapter for the existing Dataset summary presentation."""
    index = index_of(dataset)
    table = index.frame
    total, labeled = len(index), int(table.identity_id.notna().sum())
    split_statistics = {}
    for split in index.splits:
        rows = table.loc[table.split == split]
        count = int(rows.identity_id.notna().sum())
        split_statistics[split] = dict(images=len(rows), annotated=count, annotations=count, empty=len(rows)-count)
    counts = table.identity_id.value_counts()
    classes = {key: dict(name=label, images=int(counts.get(label, 0)), annotations=int(counts.get(label, 0)),
                        image_fraction=float(counts.get(label, 0))/total if total else 0.) for key, label in dataset.classes.items()}
    return dict(splits={split: stats['images'] for split, stats in split_statistics.items()}, images=total,
                annotations=labeled, annotated=labeled, empty=total-labeled, empty_fraction=(total-labeled)/total if total else 0.,
                image_sizes=[tuple(map(int, row)) for row in table[['width','height']].drop_duplicates().itertuples(index=False, name=None)],
                split_statistics=split_statistics, class_statistics=classes, mask_statistics=None)


def stats(data, *, by="identity_id"):
    """Count identity observations using metadata only, returning a pandas table.

    Parameters:
        data: Identity Dataset.
        by: Metadata field or tuple of fields to group by; () returns one total row.

    Missing identities remain an explicit unlabeled group. Tracklets are counted
    by (clip_id, tracklet_id), never by the number of archive parts.
    """
    index = index_of(data)
    fields = (by,) if isinstance(by, str) else tuple(by)
    allowed = {"identity_id", "clip_id", "tracklet_id", "split", "local_id"}
    if len(set(fields)) != len(fields) or set(fields)-allowed:
        raise ValueError(f"by must contain distinct metadata fields from {sorted(allowed)}")
    table = index.frame[list(allowed)].astype(object).copy()
    table["_tracklet"] = pd.factorize(pd.MultiIndex.from_frame(table[["clip_id", "tracklet_id"]]))[0]
    table.loc[table.tracklet_id.isna(), "_tracklet"] = np.nan
    table["_labeled"] = table.identity_id.notna().astype(int)
    if not fields:
        table["_total"] = "all"
    result = table.groupby(list(fields) or ["_total"], dropna=False, sort=True).agg(
        samples=("_labeled", "size"), labeled=("_labeled", "sum"),
        individuals=("identity_id", "nunique"), clips=("clip_id", "nunique"), tracklets=("_tracklet", "nunique"),
    ).reset_index()
    result["unlabeled"] = result.samples-result.labeled
    return result.drop(columns="_total", errors="ignore")


def filter_predictions(dataset, predictions, *, has_prediction, gt_annotated=None, split=None):
    """Match prediction observations, including separate rows of one NPZ archive."""
    if predictions is None or predictions.task != "id":
        raise ValueError("Identity filtering needs predictions=model.predict(..., gallery=gallery)")
    index = index_of(dataset)
    if split is not None:
        index = index.where(split=split)
    columns = ["path", "array", "row", "crop", "trim_black_borders"]
    def keys(frame):
        table = frame[columns].astype(object).copy()
        table['row'] = table['row'].fillna(-1).astype('int64')
        table['array'] = table['array'].fillna('')
        table['trim_black_borders'] = table['trim_black_borders'].fillna(False).astype(bool)
        table['crop'] = table['crop'].map(lambda value: str(value) if isinstance(value, (list, tuple)) else '')
        return pd.util.hash_pandas_object(table, index=False)
    refs = []
    for record in predictions:
        ref = json.loads(record.metadata['source_ref'])
        refs.append(dict(zip(columns, [ref[0], ref[1], ref[2], tuple(ref[3]) if ref[3] else None,
                                      ref[4] if len(ref)>4 else False])))
    if not refs:
        return as_dataset(index.take([]), parent=dataset)
    expected = dict(zip(keys(pd.DataFrame(refs)), [bool(record.count) for record in predictions]))
    presence = keys(index.frame).map(expected)
    mask = presence.notna() & (presence == has_prediction)
    if gt_annotated is not None:
        mask &= index.frame.identity_id.notna() == gt_annotated
    return as_dataset(index.take(np.flatnonzero(mask)), parent=dataset)
